import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"

ws['A1'] = 10
ws['B1'] = 20
ws['C1'] = "=A1+B1" # Should be 30
ws['A2'] = 5
ws['B2'] = 10
ws['C2'] = "=A2*B2" # Should be 50

wb.save('test_workbook.xlsx')
print("Created test_workbook.xlsx")
